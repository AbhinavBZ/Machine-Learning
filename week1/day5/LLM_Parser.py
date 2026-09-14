import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq

load_dotenv()
my_api_key=os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("API key kaha hai bhai")

client=Groq(api_key=my_api_key)
model="openai/gpt-oss-120b"


job_description="""
Accenture: 
Software Development InternRole Overview: 
Gain hands-on experience working on real client-facing and 
internal software solutions under the mentorship of experienced professionals.Key Responsibilities:
Contribute to the full software development lifecycle by designing and implementing custom features.
Write clean, maintainable, and testable code using modern frameworks and languages (e.g., Java, Python, C#, JavaScript).
Participate in team code reviews, debug applications, and resolve system issues to maintain reliability.
Learn and apply Accenture's enterprise standards related to software engineering, security, and compliance.
Basic Qualifications:Currently pursuing an undergraduate or master's degree in Computer Science, Software Engineering, Information Technology, or a related discipline.
Understanding of fundamental concepts like object-oriented programming, data structures, and algorithms.Strong analytical thinking, problem-solving skills, and a proactive willingness to learn.
"""

from pydantic import BaseModel
class JobD(BaseModel):
    role:str
    required_skills:list[str]
    preferred_skills:list[str]
    minimum_experience:float|None
    education_requirements:list[str]
    responsibilities:list[str]

jobd_schema=JobD.model_json_schema()

system_prompt=f"""
You are an expert HR assistant.
Your job is to analyze job descriptions and extract structured information from them.
Resturn ONLY Valid JSON matching this schema:
{jobd_schema}

IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "Properties","title" or "type".
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing , return an empty list.
Do not invent information.
"""

user_prompt=f"""

Analyze the following job description:
{job_description}
"""
message_system={
    "role":"system",
    "content":system_prompt
}
message_user={
    "role":"user",
    "content":user_prompt
}
response_format={
    "type":"json_object"
}


messages=[message_system,message_user]

response=client.chat.completions.create(model=model,messages=messages,response_format=response_format)
answer=response.choices[0].message.content

raw_json=answer

#print(raw_json)


#---------------------------------------------PART 2

import json
jobd_data=json.loads(raw_json)

job=JobD(**jobd_data)

print(job.minimum_experience)
print(job.education_requirements)

#parse real
class MatchResult(BaseModel):
    score:float
    details:dict
class Experience(BaseModel):
    company:str|None=None
    role:str|None=None
    duration:str|None=None
    description:str|None=None
    skills_used:list[str]=[]
class Resume(BaseModel):
    name:str|None=None
    email:str|None=None
    phone:str|None=None

    total_experience_years:float|None=None

    skills:list[str]=[]
    experiences:list[Experience]=[]
    education:list[str]=[]
    projects:list[str]=[]
    certifications:list[str]=[]

resume_schema=Resume.model_json_schema()
def final_score(job,resume):
    match_schema=MatchResult.model_json_schema()
    prompt=f"""
    You are an HR recruiter.
    Compare the candidates's resume with the job description.
    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:
    {match_schema}

    Give me:
    1. Candidate Name
    2. Matching Skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.

    """

    message={
        "role":"user",
        "content":prompt
    }
    messages=[message]
    response_format={
        "type":"json_object"
    }
    response=client.chat.completions.create(model=model,messages=messages,response_format=response_format)
    data=json.loads(response.choices[0].message.content)
    return MatchResult(**data)
def parse_resume(resume_text):
    system_prompt=f"""
    You are an expert resume parser.
    Extract information from the resume based on its meaning.
    not only based on extact section headings.
    Different resumes may use different headings. 

    For example:
    -Experience
    -Professional Experience
    -Work History
    -Employment
    -Internships

    These may all contain relevet experience.

    Skills may also appear in the skills section, work experience, internships or projects.

    Return ONLY valid JSON matching this schema:
    {resume_schema}

    Important rules:
    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """

    user_prompt=f"""
    Parse the following resume:
    {resume_text}
    """
    message_system={
        "role":"system",
        "content":system_prompt
    }
    message_user={
        "role":"user",
        "content":user_prompt
    }
    messages=[message_system,message_user]
    response_format={
        "type":"json_object"
    }
    response=client.chat.completions.create(model=model,messages=messages,response_format=response_format)
    raw_output=response.choices[0].message.content
    data=json.loads(raw_output)
    resume=Resume(**data)
    return resume

from pypdf import PdfReader
from docx import Document
import time

def read_pdf(file_path):
    reader=PdfReader(file_path)
    text=""
    for page in reader.pages:
        page_text=page.extract_text()
        if page_text:
            text+=page_text+"\n"
    return text
def read_docx(file_path):
    document=Document(file_path)
    text=""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text+=paragraph.text+"\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text+=cell.text+"\n"
    return text
def read_resume(file_path):
    if file_path.suffix.lower()==".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower()==".docx":
        return read_docx(file_path)
    else:
        return None


#lets do it now
resume_folder=Path("resumes")
all_results=[]
for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf",".docx"]:
        continue
    print("\nProcessing:", file_path.name)
    resume_text=read_resume(file_path)
    parsed_resume=parse_resume(resume_text)
    time.sleep(5)
    result=final_score(job,parsed_resume)
    time.sleep(5)
    print("Score:",result.score)
    all_results.append({
        "name":parsed_resume.name,
        "score":result.score,
        "details":result.details
    })
all_results.sort(
    key=lambda candidate:candidate["score"],
    reverse=True
)
top_2=all_results[:2]
worst_2=all_results[-2:]

print("Top 2 Candidates : ")
for candidate in top_2:
    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )

    print(candidate["details"])
print("Lowest 2 Candidates : ")
for candidate in worst_2:
    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )

    print(candidate["details"])


                    


